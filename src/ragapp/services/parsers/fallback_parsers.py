"""
src/ragapp/services/parsers/fallback_parsers.py

Fallback parsers for document types not supported by AWS Textract or for
simple parsing without OCR: TXT, basic DOCX, HTML, and Markdown.
"""

import re
from pathlib import Path
from typing import Optional, Any
from datetime import datetime

from ragapp.schemas.parsing import (
    ParsedDocument,
    DocumentType,
    ParsingBackend,
    Page,
    DocumentElement,
    BlockType,
)
from ragapp.services.parsers.base import SimpleParser
from ragapp.services.parsers.utils import (
    validate_file,
    validate_and_decode_text_file,
    normalize_to_utf8,
    is_html_like,
    is_markdown_like,
)
from ragapp.exception.parsing_exception import (
    UnsupportedDocumentTypeError,
    CorruptedDocumentError,
    EncodingError,
)
from ragapp.logger.custom_logger import CustomLogger

logger = CustomLogger().get_logger(__name__)


class PlainTextParser(SimpleParser):
    """Parser for plain text (.txt) files."""

    def __init__(self):
        super().__init__("Plain Text Parser")
        self._supported_types = [DocumentType.TXT]

    def supports_document_type(self, document_type: DocumentType) -> bool:
        return document_type in self._supported_types

    async def validate_file(
        self,
        file_path: str | Path,
        max_size_mb: int = 50,
    ) -> tuple[bool, Optional[str]]:
        """Validate text file."""
        return await validate_file(file_path, max_size_mb)

    async def parse(
        self,
        file_path: str | Path,
        document_type: DocumentType,
        metadata: Optional[dict[str, Any]] = None,
    ) -> ParsedDocument:
        """Parse plain text file."""
        file_path = Path(file_path)
        
        logger.info("Starting plain text parsing", file=str(file_path))
        
        try:
            # Decode text file
            content, encoding = validate_and_decode_text_file(file_path)
            
            # Split into paragraphs/sections
            paragraphs = content.split("\n\n")
            
            # Create elements
            elements: list[DocumentElement] = []
            for i, para in enumerate(paragraphs):
                if para.strip():
                    element = DocumentElement(
                        element_id=f"para_{i}",
                        type=BlockType.TEXT,
                        page_number=1,
                        content=para.strip(),
                    )
                    elements.append(element)
            
            # Create single page
            page = Page(
                page_number=1,
                text=content,
                element_count=len(elements),
                elements=elements,
            )
            
            # Create document
            document = ParsedDocument(
                document_id=f"{file_path.stem}_{int(datetime.now().timestamp())}",
                file_name=file_path.name,
                document_type=DocumentType.TXT,
                parsing_backend=ParsingBackend.SIMPLE_PARSER,
                full_text=content,
                pages=[page],
                total_pages=1,
                total_tables=0,
                total_forms=0,
                encoding=encoding,
                custom_metadata=metadata or {},
            )
            
            logger.info("Successfully parsed plain text file",
                       file=str(file_path), elements=len(elements))
            
            return document
            
        except EncodingError as e:
            logger.error("Encoding error in text file", file=str(file_path))
            raise
        except Exception as e:
            logger.error("Error parsing text file",
                        file=str(file_path), error=str(e))
            raise CorruptedDocumentError(file_path.name, str(e), e)

    def detect_encoding(self, file_path: str | Path) -> str:
        """Detect file encoding."""
        from ragapp.services.parsers.utils import detect_encoding
        return detect_encoding(file_path)

    def normalize_encoding(self, content: str, target_encoding: str = "utf-8") -> str:
        """Normalize text encoding."""
        return normalize_to_utf8(content, target_encoding)


class MarkdownParser(SimpleParser):
    """Parser for Markdown files (.md, .markdown)."""

    def __init__(self):
        super().__init__("Markdown Parser")
        self._supported_types = [DocumentType.MARKDOWN]

    def supports_document_type(self, document_type: DocumentType) -> bool:
        return document_type in self._supported_types

    async def validate_file(
        self,
        file_path: str | Path,
        max_size_mb: int = 50,
    ) -> tuple[bool, Optional[str]]:
        """Validate markdown file."""
        return await validate_file(file_path, max_size_mb)

    async def parse(
        self,
        file_path: str | Path,
        document_type: DocumentType,
        metadata: Optional[dict[str, Any]] = None,
    ) -> ParsedDocument:
        """Parse Markdown file with structure preservation."""
        file_path = Path(file_path)
        
        logger.info("Starting Markdown parsing", file=str(file_path))
        
        try:
            # Decode markdown file
            content, encoding = validate_and_decode_text_file(file_path)
            
            # Parse markdown structure
            elements = self._parse_markdown_structure(content)
            
            # Create single page
            page = Page(
                page_number=1,
                text=content,
                element_count=len(elements),
                elements=elements,
            )
            
            # Create document
            document = ParsedDocument(
                document_id=f"{file_path.stem}_{int(datetime.now().timestamp())}",
                file_name=file_path.name,
                document_type=DocumentType.MARKDOWN,
                parsing_backend=ParsingBackend.MARKDOWN,
                full_text=content,
                pages=[page],
                total_pages=1,
                total_tables=0,
                total_forms=0,
                encoding=encoding,
                custom_metadata=metadata or {},
            )
            
            logger.info("Successfully parsed Markdown file",
                       file=str(file_path), elements=len(elements))
            
            return document
            
        except Exception as e:
            logger.error("Error parsing Markdown file",
                        file=str(file_path), error=str(e))
            raise CorruptedDocumentError(file_path.name, str(e), e)

    def _parse_markdown_structure(self, content: str) -> list[DocumentElement]:
        """Parse Markdown content and extract structured elements."""
        elements: list[DocumentElement] = []
        lines = content.split("\n")
        
        element_id = 0
        i = 0
        while i < len(lines):
            line = lines[i]
            
            # Heading detection
            heading_match = re.match(r"^(#+)\s+(.+)$", line)
            if heading_match:
                level = len(heading_match.group(1))
                heading_text = heading_match.group(2)
                element = DocumentElement(
                    element_id=f"heading_{element_id}",
                    type=BlockType.HEADING,
                    page_number=1,
                    content=heading_text,
                )
                elements.append(element)
                element_id += 1
                i += 1
                continue
            
            # List item detection
            list_match = re.match(r"^[-*]\s+(.+)$", line)
            if list_match:
                list_items = []
                while i < len(lines) and re.match(r"^[-*]\s+", lines[i]):
                    item_text = re.match(r"^[-*]\s+(.+)$", lines[i]).group(1)
                    list_items.append(item_text)
                    i += 1
                
                element = DocumentElement(
                    element_id=f"list_{element_id}",
                    type=BlockType.LIST,
                    page_number=1,
                    content="\n".join(list_items),
                )
                elements.append(element)
                element_id += 1
                continue
            
            # Code block detection
            if line.strip().startswith("```"):
                code_lines = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    code_lines.append(lines[i])
                    i += 1
                i += 1  # Skip closing ```
                
                element = DocumentElement(
                    element_id=f"code_{element_id}",
                    type=BlockType.CODE,
                    page_number=1,
                    content="\n".join(code_lines),
                )
                elements.append(element)
                element_id += 1
                continue
            
            # Regular paragraph
            if line.strip():
                element = DocumentElement(
                    element_id=f"text_{element_id}",
                    type=BlockType.TEXT,
                    page_number=1,
                    content=line.strip(),
                )
                elements.append(element)
                element_id += 1
            
            i += 1
        
        return elements

    def detect_encoding(self, file_path: str | Path) -> str:
        """Detect file encoding."""
        from ragapp.services.parsers.utils import detect_encoding
        return detect_encoding(file_path)

    def normalize_encoding(self, content: str, target_encoding: str = "utf-8") -> str:
        """Normalize text encoding."""
        return normalize_to_utf8(content, target_encoding)


class HTMLParser(SimpleParser):
    """Parser for HTML files (.html, .htm)."""

    def __init__(self):
        super().__init__("HTML Parser")
        self._supported_types = [DocumentType.HTML]

    def supports_document_type(self, document_type: DocumentType) -> bool:
        return document_type in self._supported_types

    async def validate_file(
        self,
        file_path: str | Path,
        max_size_mb: int = 50,
    ) -> tuple[bool, Optional[str]]:
        """Validate HTML file."""
        return await validate_file(file_path, max_size_mb)

    async def parse(
        self,
        file_path: str | Path,
        document_type: DocumentType,
        metadata: Optional[dict[str, Any]] = None,
    ) -> ParsedDocument:
        """Parse HTML file and extract text."""
        file_path = Path(file_path)
        
        logger.info("Starting HTML parsing", file=str(file_path))
        
        try:
            # Try BeautifulSoup first
            try:
                from bs4 import BeautifulSoup
                
                with open(file_path, "rb") as f:
                    soup = BeautifulSoup(f, "html.parser")
                
                # Extract text
                text = soup.get_text(separator="\n", strip=True)
                encoding = "utf-8"
                
            except ImportError:
                # Fallback to regex-based parsing
                logger.info("BeautifulSoup not available, using regex parsing")
                content, encoding = validate_and_decode_text_file(file_path)
                text = self._extract_text_from_html_regex(content)
            
            # Normalize encoding
            text = normalize_to_utf8(text, encoding)
            
            # Create elements
            elements = self._parse_html_structure(text)
            
            # Create single page
            page = Page(
                page_number=1,
                text=text,
                element_count=len(elements),
                elements=elements,
            )
            
            # Create document
            document = ParsedDocument(
                document_id=f"{file_path.stem}_{int(datetime.now().timestamp())}",
                file_name=file_path.name,
                document_type=DocumentType.HTML,
                parsing_backend=ParsingBackend.BEAUTIFULSOUP,
                full_text=text,
                pages=[page],
                total_pages=1,
                total_tables=0,
                total_forms=0,
                encoding=encoding,
                custom_metadata=metadata or {},
            )
            
            logger.info("Successfully parsed HTML file",
                       file=str(file_path), elements=len(elements))
            
            return document
            
        except Exception as e:
            logger.error("Error parsing HTML file",
                        file=str(file_path), error=str(e))
            raise CorruptedDocumentError(file_path.name, str(e), e)

    def _extract_text_from_html_regex(self, html_content: str) -> str:
        """Extract text from HTML using regex (fallback)."""
        # Remove script and style tags
        text = re.sub(r"<script[^>]*>.*?</script>", "", html_content, flags=re.DOTALL)
        text = re.sub(r"<style[^>]*>.*?</style>", "", text, flags=re.DOTALL)
        
        # Remove HTML tags
        text = re.sub(r"<[^>]+>", "\n", text)
        
        # Decode HTML entities
        import html
        text = html.unescape(text)
        
        # Clean up whitespace
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        return "\n".join(lines)

    def _parse_html_structure(self, content: str) -> list[DocumentElement]:
        """Parse HTML content and extract elements."""
        elements: list[DocumentElement] = []
        
        lines = content.split("\n")
        element_id = 0
        
        for line in lines:
            if line.strip():
                # Simple heuristic: if line is short and ends with ":", it's likely a heading
                if len(line) < 100 and line.strip().endswith(":"):
                    element = DocumentElement(
                        element_id=f"heading_{element_id}",
                        type=BlockType.HEADING,
                        page_number=1,
                        content=line.strip(),
                    )
                else:
                    element = DocumentElement(
                        element_id=f"text_{element_id}",
                        type=BlockType.TEXT,
                        page_number=1,
                        content=line.strip(),
                    )
                
                elements.append(element)
                element_id += 1
        
        return elements

    def detect_encoding(self, file_path: str | Path) -> str:
        """Detect file encoding."""
        from ragapp.services.parsers.utils import detect_encoding
        return detect_encoding(file_path)

    def normalize_encoding(self, content: str, target_encoding: str = "utf-8") -> str:
        """Normalize text encoding."""
        return normalize_to_utf8(content, target_encoding)


class SimpleDOCXParser(SimpleParser):
    """Simple parser for DOCX files without full table support."""

    def __init__(self):
        super().__init__("Simple DOCX Parser")
        self._supported_types = [DocumentType.DOCX]

    def supports_document_type(self, document_type: DocumentType) -> bool:
        return document_type in self._supported_types

    async def validate_file(
        self,
        file_path: str | Path,
        max_size_mb: int = 50,
    ) -> tuple[bool, Optional[str]]:
        """Validate DOCX file."""
        return await validate_file(file_path, max_size_mb)

    async def parse(
        self,
        file_path: str | Path,
        document_type: DocumentType,
        metadata: Optional[dict[str, Any]] = None,
    ) -> ParsedDocument:
        """Parse DOCX file."""
        file_path = Path(file_path)
        
        logger.info("Starting simple DOCX parsing", file=str(file_path))
        
        try:
            from docx import Document
            
            # Load document
            doc = Document(file_path)
            
            # Extract text
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            
            full_text = "\n".join(paragraphs)
            
            # Create elements
            elements: list[DocumentElement] = []
            for i, para_text in enumerate(paragraphs):
                element = DocumentElement(
                    element_id=f"para_{i}",
                    type=BlockType.TEXT,
                    page_number=1,
                    content=para_text,
                )
                elements.append(element)
            
            # Create single page
            page = Page(
                page_number=1,
                text=full_text,
                element_count=len(elements),
                elements=elements,
            )
            
            # Create document
            document = ParsedDocument(
                document_id=f"{file_path.stem}_{int(datetime.now().timestamp())}",
                file_name=file_path.name,
                document_type=DocumentType.DOCX,
                parsing_backend=ParsingBackend.PYTHON_DOCX,
                full_text=full_text,
                pages=[page],
                total_pages=1,
                total_tables=0,
                total_forms=0,
                custom_metadata=metadata or {},
            )
            
            logger.info("Successfully parsed DOCX file",
                       file=str(file_path), paragraphs=len(paragraphs))
            
            return document
            
        except ImportError:
            logger.error("python-docx not installed, required for DOCX parsing")
            raise CorruptedDocumentError(
                file_path.name,
                "python-docx library not installed"
            )
        except Exception as e:
            logger.error("Error parsing DOCX file",
                        file=str(file_path), error=str(e))
            raise CorruptedDocumentError(file_path.name, str(e), e)

    def detect_encoding(self, file_path: str | Path) -> str:
        """DOCX files use internal XML encoding (usually UTF-8)."""
        return "utf-8"

    def normalize_encoding(self, content: str, target_encoding: str = "utf-8") -> str:
        """Normalize text encoding."""
        return normalize_to_utf8(content, target_encoding)
